from pathlib import Path
import re
p=Path('main.py')
s=p.read_text(encoding='utf-8')
s=s.replace("APP_VERSION = '6.0.3'", "APP_VERSION = '6.0.4'")
needle="""try:\n    from docx import Document\nexcept Exception:\n    Document = None\n\nfrom face_engine import FaceEngine\n"""
repl="""try:\n    from docx import Document\nexcept Exception:\n    Document = None\n\ntry:\n    from reportlab.lib.pagesizes import A4\n    from reportlab.lib import colors\n    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle\n    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak\n    from reportlab.pdfbase import pdfmetrics\n    from reportlab.pdfbase.ttfonts import TTFont\nexcept Exception:\n    A4 = colors = getSampleStyleSheet = ParagraphStyle = SimpleDocTemplate = Paragraph = Spacer = Table = TableStyle = RLImage = PageBreak = None\n    pdfmetrics = TTFont = None\n\nfrom face_engine import FaceEngine\n"""
s=s.replace(needle,repl)
s=s.replace("'thumb_size':72}", "'thumb_size':72,'sort_field':'archive_no','sort_dir':'asc'}")
old="""        ttk.Button(bar,text='Каталог 🔒',command=lambda:self.app.show_screen('catalog')).pack(side='left',padx=3)\n        ttk.Button(bar,text='Заблокировать',command=self.app.reset_lock).pack(side='right',padx=3)\n"""
new="""        ttk.Button(bar,text='Каталог 🔒',command=lambda:self.app.show_screen('catalog')).pack(side='left',padx=3)\n        ttk.Label(bar,text=f'Версия {APP_VERSION}',font=('Arial',10,'bold')).pack(side='right',padx=(10,4))\n        ttk.Button(bar,text='Заблокировать',command=self.app.reset_lock).pack(side='right',padx=3)\n"""
s=s.replace(old,new)
marker='class PhotoCatalogApp(tk.Tk):'
sort_class=r'''class SortDialog(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app)
        self.app=app; self.title('Сортировка документов'); self.resizable(False,False); self.transient(app); self.grab_set()
        frm=ttk.Frame(self,padding=14); frm.pack(fill='both',expand=True)
        ttk.Label(frm,text='Упорядочить документы:',font=('Arial',11,'bold')).grid(row=0,column=0,columnspan=2,sticky='w',pady=(0,10))
        self.field=tk.StringVar(value=app.settings.get('sort_field','archive_no'))
        self.direction=tk.StringVar(value=app.settings.get('sort_dir','asc'))
        ttk.Radiobutton(frm,text='По архивному номеру',variable=self.field,value='archive_no').grid(row=1,column=0,columnspan=2,sticky='w',pady=3)
        ttk.Radiobutton(frm,text='По дате съёмки',variable=self.field,value='shot_date').grid(row=2,column=0,columnspan=2,sticky='w',pady=3)
        ttk.Separator(frm).grid(row=3,column=0,columnspan=2,sticky='ew',pady=10)
        ttk.Radiobutton(frm,text='По возрастанию / старые → новые',variable=self.direction,value='asc').grid(row=4,column=0,columnspan=2,sticky='w',pady=3)
        ttk.Radiobutton(frm,text='По убыванию / новые → старые',variable=self.direction,value='desc').grid(row=5,column=0,columnspan=2,sticky='w',pady=3)
        b=ttk.Frame(frm); b.grid(row=6,column=0,columnspan=2,sticky='e',pady=(12,0))
        ttk.Button(b,text='Применить',command=self.apply).pack(side='left',padx=4)
        ttk.Button(b,text='Отмена',command=self.destroy).pack(side='left')
        self.update_idletasks(); self.geometry(f'+{app.winfo_rootx()+180}+{app.winfo_rooty()+120}')
    def apply(self):
        self.app.settings['sort_field']=self.field.get(); self.app.settings['sort_dir']=self.direction.get(); self.app.save_settings()
        for name in ('viewer','editor','catalog'):
            try:self.app.screens[name].refresh()
            except Exception:pass
        self.destroy()


'''
s=s.replace(marker, sort_class+marker)
s=s.replace("    def open_settings(self):\n        SettingsDialog(self)\n", "    def open_settings(self):\n        SettingsDialog(self)\n\n    def open_sorting(self):\n        SortDialog(self)\n")
old="""    def refresh(self): pass\n\n\nclass ViewerScreen"""
new=r'''    def sort_rows(self, rows):
        rows=list(rows)
        field=self.app.settings.get('sort_field','archive_no')
        reverse=self.app.settings.get('sort_dir','asc')=='desc'
        def archive_key(r):
            text=safe_text(r['archive_no']).casefold()
            return tuple(int(x) if x.isdigit() else x for x in re.split(r'(\d+)', text))
        def date_key(r):
            text=safe_text(r['shot_date']).strip()
            for fmt in ('%d.%m.%Y','%Y-%m-%d','%d/%m/%Y','%Y.%m.%d','%d-%m-%Y'):
                try:return (0, datetime.strptime(text,fmt))
                except Exception:pass
            m=re.search(r'(19|20)\d{2}',text)
            if m:
                try:return (1, datetime(int(m.group()),1,1))
                except Exception:pass
            return (2,text.casefold())
        key=date_key if field=='shot_date' else archive_key
        try:return sorted(rows,key=key,reverse=reverse)
        except Exception:return rows
    def refresh(self): pass


class ViewerScreen'''
s=s.replace(old,new)
s=s.replace("""        ttk.Button(search,text='🖨 Печать результата',command=self.print_results).pack(fill='x',pady=(6,0))\n""", """        ttk.Button(search,text='🖨 Печать результата',command=self.print_results).pack(fill='x',pady=(6,0))\n        ex=ttk.Frame(search); ex.pack(fill='x',pady=(6,0))\n        ttk.Button(ex,text='Экспорт в Word',command=self.export_word).pack(side='left',fill='x',expand=True)\n        ttk.Button(ex,text='Экспорт в PDF',command=self.export_pdf).pack(side='left',fill='x',expand=True,padx=(5,0))\n        ttk.Button(search,text='↕ Сортировка',command=self.app.open_sorting).pack(fill='x',pady=(6,0))\n""")
s=s.replace("""        self.preview=ttk.Label(right,text='Выберите запись',anchor='center'); self.preview.pack(fill='both',expand=True)\n        self.preview_ref=None; self.info=tk.Text(right,height=20,wrap='word',state='disabled',font=('Arial',12),spacing1=3,spacing3=4); self.info.pack(fill='both',expand=True,pady=(8,0))\n""", """        self.preview=ttk.Label(right,text='Выберите запись',anchor='center'); self.preview.pack(fill='both',expand=True)\n        self.preview_ref=None\n        info_frame=ttk.Frame(right); info_frame.pack(fill='both',expand=True,pady=(8,0))\n        self.info=tk.Text(info_frame,height=20,wrap='word',state='disabled',font=('Arial',12),spacing1=3,spacing3=4)\n        info_scroll=ttk.Scrollbar(info_frame,orient='vertical',command=self.info.yview); self.info.configure(yscrollcommand=info_scroll.set)\n        self.info.pack(side='left',fill='both',expand=True); info_scroll.pack(side='right',fill='y')\n""")
s=s.replace("self.rows=self.db.search(self.q.get(),self.field.get(),self.whole.get());", "self.rows=self.sort_rows(self.db.search(self.q.get(),self.field.get(),self.whole.get()));")
s=s.replace("self.rows=[self.db.get(pid) for pid in ids]", "self.rows=self.sort_rows([self.db.get(pid) for pid in ids])")
needle="    def search_face(self):\n"
export_methods=r'''    def export_word(self):
        if not self.rows:
            messagebox.showinfo(APP_NAME,'Нет результатов для экспорта.'); return
        if Document is None:
            messagebox.showerror(APP_NAME,'Модуль Microsoft Word DOCX недоступен.'); return
        path=filedialog.asksaveasfilename(title='Сохранить результаты в Word',defaultextension='.docx',filetypes=[('Microsoft Word','*.docx')],initialfile=f'Результаты_поиска_{datetime.now():%Y%m%d_%H%M}.docx')
        if not path:return
        try:
            from docx.shared import Inches
            doc=Document(); doc.add_heading('Архив Президента Кыргызской Республики — Фотокаталог',level=1)
            doc.add_paragraph(f'Найдено записей: {len(self.rows)}')
            for n,r in enumerate(self.rows,1):
                doc.add_heading(f"{n}. Архивный номер: {safe_text(r['archive_no'])}",level=2)
                fp=safe_text(r['file_path'])
                if fp and os.path.exists(fp):
                    try:doc.add_picture(fp,width=Inches(4.8))
                    except Exception:pass
                table=doc.add_table(rows=0,cols=2); table.style='Table Grid'
                for label,key in [('Дата съёмки','shot_date'),('Место съёмки','location'),('Автор','author'),('Источник поступления','source'),('Описание','description')]:
                    cells=table.add_row().cells; cells[0].text=label; cells[1].text=safe_text(r[key])
                doc.add_paragraph('')
            doc.save(path); messagebox.showinfo(APP_NAME,f'Файл Word сохранён:\n{path}')
        except Exception as exc:messagebox.showerror(APP_NAME,f'Не удалось создать Word-файл:\n{exc}')

    def export_pdf(self):
        if not self.rows:
            messagebox.showinfo(APP_NAME,'Нет результатов для экспорта.'); return
        if SimpleDocTemplate is None:
            messagebox.showerror(APP_NAME,'Модуль создания PDF недоступен.'); return
        path=filedialog.asksaveasfilename(title='Сохранить результаты в PDF',defaultextension='.pdf',filetypes=[('PDF','*.pdf')],initialfile=f'Результаты_поиска_{datetime.now():%Y%m%d_%H%M}.pdf')
        if not path:return
        try:
            font='Helvetica'; candidates=[Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'/'arial.ttf', Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
            for fpfont in candidates:
                if fpfont.exists(): pdfmetrics.registerFont(TTFont('ArchiveArial',str(fpfont))); font='ArchiveArial'; break
            styles=getSampleStyleSheet(); normal=ParagraphStyle('ArchiveNormal',parent=styles['Normal'],fontName=font,fontSize=9,leading=12); title=ParagraphStyle('ArchiveTitle',parent=styles['Heading1'],fontName=font,fontSize=15,leading=18)
            doc=SimpleDocTemplate(path,pagesize=A4,rightMargin=30,leftMargin=30,topMargin=32,bottomMargin=32)
            story=[Paragraph('Архив Президента Кыргызской Республики — Фотокаталог',title),Spacer(1,8),Paragraph(f'Найдено записей: {len(self.rows)}',normal),Spacer(1,10)]
            for n,r in enumerate(self.rows,1):
                story.append(Paragraph(f"{n}. Архивный номер: {html.escape(safe_text(r['archive_no']))}",normal)); story.append(Spacer(1,5)); fp=safe_text(r['file_path'])
                if fp and os.path.exists(fp):
                    try: im=RLImage(fp); im._restrictSize(360,250); story.extend([im,Spacer(1,5)])
                    except Exception:pass
                data=[]
                for label,key in [('Дата съёмки','shot_date'),('Место съёмки','location'),('Автор','author'),('Источник поступления','source'),('Описание','description')]:
                    data.append([Paragraph(label,normal),Paragraph(html.escape(safe_text(r[key])).replace('\n','<br/>'),normal)])
                tbl=Table(data,colWidths=[120,390]); tbl.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.5,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP')])); story.extend([tbl,Spacer(1,12)])
            doc.build(story); messagebox.showinfo(APP_NAME,f'PDF сохранён:\n{path}')
        except Exception as exc:messagebox.showerror(APP_NAME,f'Не удалось создать PDF:\n{exc}')

'''
s=s.replace(needle,export_methods+needle,1)
s=s.replace("super().__init__(parent,app);self.nav('editor');self.current_id=None;self.current_path='';self.preview_ref=None", "super().__init__(parent,app);self.nav('editor');self.current_id=None;self.current_path='';self.preview_ref=None;self._loading=False;self._autosave_job=None")
s=s.replace("""        ttk.Label(form,text='Описание фотографии').grid(row=row,column=0,sticky='w');row+=1;self.desc=tk.Text(form,height=6,wrap='word');self.desc.grid(row=row,column=0,sticky='nsew',pady=(0,6));row+=1\n""", """        ttk.Label(form,text='Описание фотографии').grid(row=row,column=0,sticky='w');row+=1\n        desc_frame=ttk.Frame(form); desc_frame.grid(row=row,column=0,sticky='nsew',pady=(0,6))\n        self.desc=tk.Text(desc_frame,height=6,wrap='word'); desc_scroll=ttk.Scrollbar(desc_frame,orient='vertical',command=self.desc.yview); self.desc.configure(yscrollcommand=desc_scroll.set)\n        self.desc.pack(side='left',fill='both',expand=True); desc_scroll.pack(side='right',fill='y');row+=1\n""")
s=s.replace("""        form.columnconfigure(0,weight=1);form.rowconfigure(row-1,weight=1);self.status()\n""", """        form.columnconfigure(0,weight=1);form.rowconfigure(row-1,weight=1)\n        for e in self.entries.values(): e.bind('<KeyRelease>',self.schedule_autosave); e.bind('<FocusOut>',self.schedule_autosave)\n        self.desc.bind('<<Modified>>',self._desc_modified)\n        self.status()\n""",1)
s=s.replace("self.rows=self.db.search(self.q.get(),'Все поля',True)", "self.rows=self.sort_rows(self.db.search(self.q.get(),'Все поля',True))",1)
old="""        r=self.db.get(int(s[0]));self.current_id=r['id'];self.current_path=r['file_path'] or ''\n        for k,e in self.entries.items():e.delete(0,'end');e.insert(0,r[k] or '')\n        self.desc.delete('1.0','end');self.desc.insert('1.0',r['description'] or '');self.show_preview()\n"""
new="""        r=self.db.get(int(s[0]));self.current_id=r['id'];self.current_path=r['file_path'] or '';self._loading=True\n        for k,e in self.entries.items():e.delete(0,'end');e.insert(0,r[k] or '')\n        self.desc.delete('1.0','end');self.desc.insert('1.0',r['description'] or '');self.desc.edit_modified(False);self._loading=False;self.show_preview()\n"""
s=s.replace(old,new)
s=s.replace("""    def new(self):\n        self.current_id=None;self.current_path='';\n        for e in self.entries.values():e.delete(0,'end')\n        self.desc.delete('1.0','end');self.preview.configure(image='',text='Фото не выбрано');self.preview_ref=None\n""", """    def new(self):\n        self._loading=True;self.current_id=None;self.current_path='';\n        for e in self.entries.values():e.delete(0,'end')\n        self.desc.delete('1.0','end');self.desc.edit_modified(False);self.preview.configure(image='',text='Фото не выбрано');self.preview_ref=None;self._loading=False\n""")
s=s.replace("if p:self.current_path=p;self.show_preview()", "if p:self.current_path=p;self.show_preview();self.schedule_autosave()")
s=s.replace("def remove_photo(self):self.current_path='';self.show_preview()", "def remove_photo(self):self.current_path='';self.show_preview();self.schedule_autosave()")
needle="    def collect(self):\n"
auto=r'''    def _desc_modified(self,_e=None):
        if self.desc.edit_modified(): self.desc.edit_modified(False); self.schedule_autosave()
    def schedule_autosave(self,_e=None):
        if self._loading:return
        if self._autosave_job:
            try:self.after_cancel(self._autosave_job)
            except Exception:pass
        self._autosave_job=self.after(700,self._autosave)
    def _autosave(self):
        self._autosave_job=None
        if self._loading:return
        d=self.collect()
        if not d['archive_no']: self.status_var.set('Автосохранение: укажите архивный номер.'); return
        try:
            old=self.db.get(self.current_id) if self.current_id else None; p=d['file_path']
            if p and os.path.exists(p):
                if old and old['file_path']==p:d['file_path']=p
                elif str(Path(p).resolve()).startswith(str(PHOTOS_DIR.resolve())):d['file_path']=p
                else:d['file_path']=copy_into_archive(p)
            if self.current_id:self.db.update(self.current_id,d)
            else:self.current_id=self.db.add(d)
            self.current_path=d['file_path']; self.status_var.set('Изменения сохранены автоматически.')
        except Exception as exc:self.status_var.set(f'Ошибка автосохранения: {exc}')

'''
s=s.replace(needle,auto+needle,1)
s=s.replace("self.current_path=d['file_path'];self.refresh();self.status_var.set('Сохранено.')", "self.current_path=d['file_path'];self.refresh();self.status_var.set('Сохранено. Автосохранение включено.')")
s=s.replace("ttk.Button(top1,text='⚙ Вид',command=self.app.open_settings).pack(side='right',padx=(8,0))", "ttk.Button(top1,text='⚙ Вид',command=self.app.open_settings).pack(side='right',padx=(8,0));ttk.Button(top1,text='↕ Сортировка',command=self.app.open_sorting).pack(side='right',padx=(8,0))")
s=s.replace("ttk.Button(top2,text='Удалить выбранные',command=self.delete_selected).pack(side='left',padx=4)", "ttk.Button(top2,text='Удалить выбранные',command=self.delete_selected).pack(side='left',padx=4);ttk.Button(top2,text='Копировать',command=self.copy_selected).pack(side='left',padx=4)")
s=s.replace("self.tree.bind('<Control-a>',self.ctrl_a)", "self.tree.bind('<Control-a>',self.ctrl_a);self.tree.bind('<Control-c>',self.ctrl_c);self.tree.bind('<Control-C>',self.ctrl_c)")
s=s.replace("self.all_rows=self.db.search(self.q.get(),'Все поля',True)", "self.all_rows=self.sort_rows(self.db.search(self.q.get(),'Все поля',True))")
needle="    def select_page(self):\n"
copy_methods=r'''    def ctrl_c(self,_e=None):
        self.copy_selected(); return 'break'
    def copy_selected(self):
        selected=list(self.tree.selection())
        if not selected:selected=[i for i in self.checked if self.db.get(int(i))]
        if not selected:messagebox.showinfo(APP_NAME,'Сначала выделите записи через Ctrl или Shift.'); return
        order={str(r['id']):n for n,r in enumerate(self.all_rows)}; selected=sorted(set(selected),key=lambda x:order.get(str(x),10**9))
        lines=['Архивный номер\tОписание\tДата съёмки\tМесто съёмки\tАвтор\tИсточник поступления']
        for iid in selected:
            r=self.db.get(int(iid))
            if not r:continue
            vals=[safe_text(r[k]).replace('\t',' ').replace('\r',' ').replace('\n',' ') for k in ('archive_no','description','shot_date','location','author','source')]; lines.append('\t'.join(vals))
        self.clipboard_clear(); self.clipboard_append('\n'.join(lines)); self.update_idletasks(); self.status_var.set(f'Скопировано записей: {len(lines)-1}. Можно вставить в Word или Excel (Ctrl+V).')

'''
s=s.replace(needle,copy_methods+needle,1)
p.write_text(s,encoding='utf-8')
req=Path('requirements.txt'); t=req.read_text(encoding='utf-8');
if 'reportlab==' not in t: req.write_text(t.rstrip()+'\nreportlab==4.2.5\n',encoding='utf-8')
Path('version.txt').write_text('6.0.4\n',encoding='utf-8')
iss=Path('installer.iss'); x=iss.read_text(encoding='utf-8'); x=x.replace('#define MyAppVersion "6.0.3"','#define MyAppVersion "6.0.4"'); iss.write_text(x,encoding='utf-8')
print('v6.0.4 patch applied')
